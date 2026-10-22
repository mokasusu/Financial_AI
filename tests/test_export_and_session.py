"""
Unit & Integration Tests cho Giai đoạn 3: Chuẩn hóa, Lưu trữ và Xuất Báo cáo Tài chính.
Kiểm tra:
1. Chuẩn hóa số nguyên lớn và số âm kế toán, tránh ký hiệu khoa học.
2. Xử lý giá trị thiếu (None) và không tự ý bịa số liệu.
3. Lưu và nạp phiên đối soát không làm mất trạng thái xác minh (needs_review).
4. Xử lý an toàn khi nạp JSON sai schema hoặc bị hỏng.
5. OCR gốc không bị ghi đè bởi gợi ý Vision AI.
6. Xuất workbook Excel với đủ 4 sheet bắt buộc (Financial_Data, Verification, Ratios, Metadata).
7. Dữ liệu chưa xác minh không được xuất thành số liệu chính thức đã chốt.
8. Bảng Ratios thể hiện rõ lý do chặn tính tỷ số khi dữ liệu chưa đủ điều kiện.
9. Xuất và đọc lại JSON bảo toàn tính toàn vẹn 100%.
10. Lưu tệp an toàn (Atomic Write) vào thư mục processed.
11. Xử lý an toàn khi gặp lỗi IO hoặc tệp không tồn tại.
12. Tuyệt đối không lưu khóa bí mật hoặc token vào tệp xuất.
"""

import io
import json
import os
import tempfile
from pathlib import Path

import openpyxl
import pytest

from app.models.financial_statement import (
    FinancialItem,
    FinancialStatement,
    ReconciliationSession,
)
from app.utils.export_utils import (
    build_reconciliation_session,
    export_reconciliation_to_excel,
    export_reconciliation_to_json,
    format_financial_number,
    load_reconciliation_from_json,
    load_session_from_file,
    save_session_to_processed_dir,
)


def test_format_financial_number_large_and_negative():
    """Kiểm tra hàm format_financial_number xử lý số nguyên lớn và số âm không bị ký hiệu khoa học."""
    # Số nghìn tỷ chuẩn
    assert format_financial_number(36261180908033.0) == "36,261,180,908,033"
    # Số âm
    assert format_financial_number(-150000000.0) == "-150,000,000"
    # Số lẻ thập phân
    assert format_financial_number(1234567.89) == "1,234,567.89"
    # Giá trị None
    assert format_financial_number(None) == "N/A"


def test_financial_item_provenance_and_effective_value():
    """Kiểm tra các trường metadata nguồn và hàm get_effective_value() trên FinancialItem."""
    item = FinancialItem(
        code="CURRENT_ASSETS",
        name="Tài sản ngắn hạn",
        value=36000000000000.0,
        unit="VND",
        period="2024",
        status="needs_review",
        raw_ocr="36.261.180.908.033",
        ai_suggestion=36261180908033.0,
        user_value=36261180908033.0,
        review_reason="OCR dính vết mực, đã kiểm tra theo Vision AI",
        page_number=5,
        source_line="Tài sản ngắn hạn | 100 | 36.261.180.908.033",
    )

    # Khi có user_value, get_effective_value ưu tiên user_value
    assert item.get_effective_value() == 36261180908033.0
    # OCR gốc được bảo toàn riêng biệt
    assert item.raw_ocr == "36.261.180.908.033"
    # Gợi ý AI được bảo toàn riêng biệt
    assert item.ai_suggestion == 36261180908033.0
    # Lý do đối soát được lưu vết
    assert "OCR dính vết mực" in item.review_reason


def test_save_and_load_session_preserves_needs_review_and_raw_ocr():
    """
    Kiểm tra tính bất biến của trạng thái xác minh:
    Lưu một phiên có chỉ tiêu needs_review, nạp lại -> chỉ tiêu BẮT BUỘC vẫn là needs_review.
    Tuyệt đối không tự nâng lên verified.
    """
    items = [
        FinancialItem(
            code="CURRENT_ASSETS",
            name="Tài sản ngắn hạn",
            value=1000000.0,
            status="verified",
            raw_ocr="1.000.000",
        ),
        FinancialItem(
            code="TOTAL_LIABILITIES",
            name="Tổng nợ phải trả",
            value=500000.0,
            status="needs_review",
            raw_ocr="500.000 (mờ)",
            review_reason="Chữ số mờ cần đối soát với bản gốc",
        ),
    ]

    stmt = FinancialStatement(
        company_name="Công ty Test",
        fiscal_year=2024,
        period_label="2024",
        scope="consolidated",
        status="needs_review",
        items=items,
    )

    session = build_reconciliation_session(
        source_filename="test_bctc.pdf",
        statements=[stmt],
        session_id="test_sess_001",
        ratio_status="blocked",
        ratio_block_reason="Có chỉ tiêu TOTAL_LIABILITIES ở trạng thái needs_review",
    )

    json_str = export_reconciliation_to_json(session)
    loaded_session = load_reconciliation_from_json(json_str)

    assert loaded_session.session_id == "test_sess_001"
    assert len(loaded_session.statements) == 1
    loaded_stmt = loaded_session.statements[0]
    assert loaded_stmt.status == "needs_review"

    # Kiểm tra từng chỉ tiêu
    it1 = loaded_stmt.get_item("CURRENT_ASSETS")
    assert it1 is not None and it1.status == "verified"

    it2 = loaded_stmt.get_item("TOTAL_LIABILITIES")
    assert it2 is not None
    # QUY TẮC BẮT BUỘC: Không tự chuyển sang verified
    assert it2.status == "needs_review"
    assert it2.raw_ocr == "500.000 (mờ)"
    assert it2.review_reason == "Chữ số mờ cần đối soát với bản gốc"


def test_load_corrupted_or_invalid_schema_json_raises_value_error():
    """Kiểm tra xử lý an toàn khi nạp JSON bị hỏng hoặc sai schema."""
    # 1. JSON hỏng cú pháp
    with pytest.raises(ValueError, match="JSON không đúng định dạng schema"):
        load_reconciliation_from_json("{broken_json: 123")

    # 2. JSON thiếu trường bắt buộc của ReconciliationSession
    with pytest.raises(ValueError, match="JSON không đúng định dạng schema"):
        load_reconciliation_from_json(json.dumps({"some_key": "some_val"}))


def test_ocr_raw_never_overwritten_by_ai_suggestion():
    """
    Kiểm tra nguyên tắc bảo vệ dữ liệu gốc:
    Gợi ý AI khác biệt với OCR gốc không được làm thay đổi raw_ocr.
    """
    item = FinancialItem(
        code="NET_INCOME",
        name="Lợi nhuận sau thuế",
        value=11649985224938.0,
        status="needs_review",
        raw_ocr="11.649.985.224.938 11.599.653.741.335",  # Bị dính 2 cột do OCR
        ai_suggestion=11649985224938.0,
        user_value=11649985224938.0,
    )

    session = build_reconciliation_session(
        source_filename="vinamilk_p8.pdf",
        statements=[
            FinancialStatement(
                fiscal_year=2024,
                scope="consolidated",
                status="needs_review",
                items=[item],
            )
        ],
    )

    json_str = export_reconciliation_to_json(session)
    loaded = load_reconciliation_from_json(json_str)

    loaded_item = loaded.statements[0].items[0]
    # raw_ocr vẫn giữ nguyên chuỗi OCR gốc bị dính cột
    assert loaded_item.raw_ocr == "11.649.985.224.938 11.599.653.741.335"
    assert loaded_item.ai_suggestion == 11649985224938.0
    assert loaded_item.status == "needs_review"


def test_export_excel_workbook_contains_all_four_required_sheets():
    """Kiểm tra workbook Excel sinh ra chứa đủ tối thiểu 4 sheet chuẩn hóa."""
    stmt = FinancialStatement(
        company_name="Doanh nghiệp Mẫu",
        fiscal_year=2024,
        period_label="2024",
        scope="consolidated",
        status="verified",
        items=[
            FinancialItem(code="CURRENT_ASSETS", name="Tài sản ngắn hạn", value=1500.0, status="verified"),
            FinancialItem(code="CURRENT_LIABILITIES", name="Nợ ngắn hạn", value=750.0, status="verified"),
        ],
    )

    session = build_reconciliation_session(
        source_filename="bctc_sample.pdf",
        statements=[stmt],
        session_id="excel_test_01",
        ratios={"CURRENT_RATIO": 2.0, "DEBT_EQUITY": 0.5, "ROA": None, "ROE": None},
        ratio_status="calculated",
    )

    excel_bytes = export_reconciliation_to_excel(session)
    assert len(excel_bytes) > 0

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    # BẮT BUỘC phải có đủ 4 sheet
    assert "Financial_Data" in sheet_names
    assert "Verification" in sheet_names
    assert "Ratios" in sheet_names
    assert "Metadata" in sheet_names


def test_unverified_items_not_promoted_to_official_value_in_excel():
    """
    NGUYÊN TẮC BẢO TOÀN DỮ LIỆU:
    Trong sheet Financial_Data, các chỉ tiêu chưa xác minh (status='needs_review')
    không được hiển thị số tiền ở cột 'Giá trị chính thức', mà phải đánh dấu rõ là chưa chốt.
    """
    stmt = FinancialStatement(
        company_name="Doanh nghiệp Cảnh báo",
        fiscal_year=2024,
        period_label="2024",
        scope="consolidated",
        status="needs_review",
        items=[
            FinancialItem(
                code="CURRENT_ASSETS",
                name="Tài sản ngắn hạn",
                value=2000.0,
                status="verified",
            ),
            FinancialItem(
                code="TOTAL_LIABILITIES",
                name="Tổng nợ phải trả",
                value=1200.0,
                status="needs_review",
                review_reason="Nghi vấn OCR dính chữ",
                ai_suggestion=1100.0,
            ),
        ],
    )

    session = build_reconciliation_session(
        source_filename="unverified_report.pdf",
        statements=[stmt],
        session_id="unverified_sess",
    )

    excel_bytes = export_reconciliation_to_excel(session)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws_data = wb["Financial_Data"]

    # Dòng 5 là item CURRENT_ASSETS (verified), Dòng 6 là TOTAL_LIABILITIES (needs_review)
    row_5_code = ws_data.cell(row=5, column=1).value
    row_5_val = ws_data.cell(row=5, column=4).value
    row_5_status = ws_data.cell(row=5, column=8).value
    assert row_5_code == "CURRENT_ASSETS"
    assert row_5_val == 2000.0
    assert row_5_status == "Đã xác minh"

    row_6_code = ws_data.cell(row=6, column=1).value
    row_6_val = ws_data.cell(row=6, column=4).value
    row_6_status = ws_data.cell(row=6, column=8).value
    assert row_6_code == "TOTAL_LIABILITIES"
    # Cột giá trị chính thức KHÔNG được chứa 1200.0 hay gợi ý AI 1100.0 như số liệu đã chốt
    assert row_6_val == "" or row_6_val is None
    assert "Cần đối soát" in row_6_status


def test_ratios_sheet_shows_blocked_reason_when_unverified():
    """Kiểm tra sheet Ratios ghi rõ lý do chặn khi tỷ số chưa đủ điều kiện tính."""
    stmt = FinancialStatement(
        fiscal_year=2024,
        scope="consolidated",
        status="needs_review",
        items=[],
    )

    session = build_reconciliation_session(
        source_filename="blocked_ratio.pdf",
        statements=[stmt],
        ratio_status="blocked",
        ratio_block_reason="Chỉ tiêu CURRENT_LIABILITIES chưa được xác minh đối soát.",
    )

    excel_bytes = export_reconciliation_to_excel(session)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws_ratios = wb["Ratios"]

    # Dòng 4 là CURRENT_RATIO
    val_cell = ws_ratios.cell(row=4, column=4).value
    status_cell = ws_ratios.cell(row=4, column=6).value
    note_cell = ws_ratios.cell(row=4, column=7).value

    assert val_cell == "" or val_cell is None
    assert status_cell == "Chưa thể tính"
    assert "chưa được xác minh" in note_cell


def test_roundtrip_session_json_export_and_load(tmp_path):
    """Kiểm tra tính đối xứng hai chiều: Session -> JSON -> File -> Nạp lại -> So sánh chính xác."""
    stmt = FinancialStatement(
        company_name="CTCP Kiểm thử Roundtrip",
        fiscal_year=2025,
        period_label="2025",
        scope="separate",
        status="verified",
        currency="VND",
        items=[
            FinancialItem(
                code="REVENUE",
                name="Doanh thu bán hàng",
                value=5000000000.0,
                unit="VND",
                vas_code="01",
                page_number=3,
                raw_ocr="5.000.000.000",
                status="verified",
            ),
        ],
    )

    session = build_reconciliation_session(
        source_filename="vinamilk_2025.pdf",
        statements=[stmt],
        session_id="roundtrip_999",
        source_file_hash="abc123hash",
        ratios={"CURRENT_RATIO": 1.8},
        ratio_status="calculated",
        metadata={"scanner": "Tesseract 5.5"},
    )

    # 1. Test xuất và nạp chuỗi
    json_str = export_reconciliation_to_json(session)
    loaded_1 = load_reconciliation_from_json(json_str)
    assert loaded_1.session_id == session.session_id
    assert loaded_1.statements[0].items[0].value == 5000000000.0

    # 2. Test lưu atomic ra file và nạp lại
    saved_file = save_session_to_processed_dir(session, target_dir=tmp_path)
    assert saved_file.exists()
    assert saved_file.name.startswith("session_roundtrip_999_")

    loaded_2 = load_session_from_file(saved_file)
    assert loaded_2.session_id == session.session_id
    assert loaded_2.source_file_hash == "abc123hash"
    assert loaded_2.metadata.get("scanner") == "Tesseract 5.5"


def test_atomic_save_session_handles_file_not_found(tmp_path):
    """Kiểm tra hàm load_session_from_file ném lỗi FileNotFoundError khi tệp không tồn tại."""
    with pytest.raises(FileNotFoundError):
        load_session_from_file(tmp_path / "non_existent_session.json")


def test_atomic_save_handles_io_error_cleanly(tmp_path):
    """Kiểm tra xử lý an toàn khi lưu vào đường dẫn không hợp lệ (ví dụ target_dir là một tệp)."""
    file_as_dir = tmp_path / "not_a_directory.txt"
    file_as_dir.write_text("dummy")

    session = build_reconciliation_session(
        source_filename="test.pdf",
        statements=[],
        session_id="err_sess",
    )

    with pytest.raises(IOError):
        save_session_to_processed_dir(session, target_dir=file_as_dir)


def test_session_metadata_contains_zero_api_keys():
    """Kiểm tra tuyệt đối không đưa API Key hay bí mật môi trường vào nội dung file session hoặc Excel."""
    os.environ["OPENROUTER_API_KEY"] = "sk-or-v1-secret-never-expose-this"

    stmt = FinancialStatement(
        company_name="Bảo mật Doanh nghiệp",
        fiscal_year=2024,
        scope="consolidated",
        status="verified",
        items=[],
    )

    session = build_reconciliation_session(
        source_filename="secret_check.pdf",
        statements=[stmt],
        session_id="safe_sess",
    )

    json_str = export_reconciliation_to_json(session)
    # Khóa bí mật không được xuất hiện trong JSON
    assert "sk-or-v1-secret" not in json_str

    excel_bytes = export_reconciliation_to_excel(session)
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws_meta = wb["Metadata"]

    # Quét tất cả các ô trong sheet Metadata
    for row in ws_meta.iter_rows(values_only=True):
        for cell_val in row:
            if cell_val is not None:
                assert "sk-or-v1-secret" not in str(cell_val)
