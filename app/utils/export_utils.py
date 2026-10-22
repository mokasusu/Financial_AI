"""
Module tiện ích xuất báo cáo tài chính và quản lý phiên đối soát (Giai đoạn 3).
Hỗ trợ:
- Xuất workbook Excel 4 sheet (.xlsx) bằng openpyxl có định dạng chuyên nghiệp.
- Serialize và deserialize JSON phiên đối soát (ReconciliationSession).
- Lưu trữ an toàn (atomic write) vào thư mục data/processed/ và nạp lại phiên.
- Bảo vệ dữ liệu nguồn, phân biệt rõ số liệu đã xác minh và chưa xác minh.
"""

import io
import json
import os
import re
import tempfile
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Literal

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models.financial_statement import (
    FinancialItem,
    FinancialStatement,
    ReconciliationSession,
)


def build_reconciliation_session(
    source_filename: str,
    statements: list[FinancialStatement],
    session_id: str | None = None,
    source_file_hash: str | None = None,
    company_name: str = "Doanh nghiệp BCTC",
    ratios: dict[str, Any] | None = None,
    ratio_status: Literal["calculated", "blocked", "not_attempted"] = "not_attempted",
    ratio_block_reason: str | None = None,
    metadata: dict[str, Any] | None = None,
) -> ReconciliationSession:
    """Khởi tạo hoặc cập nhật thực thể ReconciliationSession chuẩn hóa."""
    now_iso = datetime.now(timezone.utc).isoformat()
    s_id = session_id or uuid.uuid4().hex[:8]
    return ReconciliationSession(
        session_id=s_id,
        source_filename=source_filename,
        source_file_hash=source_file_hash,
        created_at=now_iso,
        updated_at=now_iso,
        company_name=company_name,
        statements=statements,
        ratios=ratios or {},
        ratio_status=ratio_status,
        ratio_block_reason=ratio_block_reason,
        metadata=metadata or {},
    )


def _safe_str(val: Any) -> str:
    """Chuyển đổi giá trị sang chuỗi an toàn."""
    if val is None:
        return ""
    return str(val).strip()


def format_financial_number(val: float | None) -> str:
    """
    Định dạng số tài chính chính xác, tránh hiển thị ký hiệu khoa học (scientific notation).
    Ví dụ: 36261180908033.0 -> "36,261,180,908,033"
    """
    if val is None:
        return "N/A"
    try:
        # Nếu là số nguyên hoặc số thực có phần thập phân bằng 0
        if isinstance(val, (int, float)) and val == int(val):
            return f"{int(val):,}"
        return f"{val:,.2f}"
    except Exception:
        return str(val)


def export_reconciliation_to_excel(session: ReconciliationSession) -> bytes:
    """
    Tạo workbook Excel đa sheet (.xlsx) chuẩn hóa bằng openpyxl.
    Bao gồm tối thiểu 4 sheet:
    1. Financial_Data: Số liệu tài chính chính thức (chỉ ghi nhận giá trị khi đã xác minh).
    2. Verification: Toàn bộ dữ liệu đối chiếu (OCR gốc, OCR parsed, gợi ý AI, trạng thái, trang).
    3. Ratios: Kết quả 4 tỷ số tài chính (hoặc ghi rõ lý do chặn tính tỷ số).
    4. Metadata: Thông tin tệp nguồn, phiên đối soát, thời gian xuất, không chứa bí mật.
    """
    wb = openpyxl.Workbook()
    # Xóa sheet mặc định ban đầu
    wb.remove(wb.active)

    # 1. Định nghĩa Palette màu và Styles
    header_fill_blue = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_fill_gray = PatternFill(start_color="595959", end_color="595959", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    normal_font = Font(name="Calibri", size=11)
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")

    fill_verified = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")  # Xanh nhạt
    fill_needs_review = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Vàng nhạt
    font_verified = Font(name="Calibri", size=10, color="276A3C", bold=True)
    font_needs_review = Font(name="Calibri", size=10, color="976100", bold=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # =========================================================================
    # SHEET 1: Financial_Data
    # =========================================================================
    ws_data = wb.create_sheet(title="Financial_Data")
    ws_data.views.sheetView[0].showGridLines = True

    # Tiêu đề Sheet 1
    ws_data["A1"] = "BẢNG CHỈ TIÊU TÀI CHÍNH CHÍNH THỨC"
    ws_data["A1"].font = title_font
    ws_data["A2"] = f"Doanh nghiệp: {session.company_name} | Nguồn: {session.source_filename}"
    ws_data["A2"].font = normal_font

    headers_data = [
        "Mã chỉ tiêu (Code)",
        "Mã VAS (VAS Code)",
        "Tên chỉ tiêu (Name)",
        "Giá trị chính thức (VND)",
        "Đơn vị tính",
        "Kỳ báo cáo",
        "Phạm vi báo cáo",
        "Trạng thái xác minh",
        "Ghi chú dữ liệu",
    ]
    ws_data.append([])
    ws_data.append(headers_data)
    header_row_idx = 4

    for col_idx in range(1, len(headers_data) + 1):
        cell = ws_data.cell(row=header_row_idx, column=col_idx)
        cell.fill = header_fill_blue
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    current_row = header_row_idx + 1
    for stmt in session.statements:
        for item in stmt.items:
            is_verified = item.status == "verified"
            # NGUYÊN TẮC BẮT BUỘC: Không đưa gợi ý AI chưa xác minh vào cột số liệu chính thức
            official_val = item.get_effective_value() if is_verified else None
            val_display = official_val if official_val is not None else ("Chưa xác minh" if not is_verified else "N/A")

            row_values = [
                item.code,
                item.vas_code or "",
                item.name,
                official_val if official_val is not None else "",
                item.unit,
                item.period or stmt.period_label or "",
                item.scope,
                "Đã xác minh" if is_verified else "Cần đối soát (Chưa chốt)",
                "" if is_verified else (item.review_reason or "Chưa được người dùng đối chiếu xác nhận"),
            ]
            ws_data.append(row_values)

            # Căn chỉnh và định dạng số
            for c_idx in range(1, len(row_values) + 1):
                c = ws_data.cell(row=current_row, column=c_idx)
                c.font = normal_font
                c.border = thin_border
                if c_idx == 4 and official_val is not None:
                    c.number_format = "#,##0"
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 8:
                    c.fill = fill_verified if is_verified else fill_needs_review
                    c.font = font_verified if is_verified else font_needs_review
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(vertical="center")

            current_row += 1

    # =========================================================================
    # SHEET 2: Verification (Đối chiếu đa nguồn OCR - AI - Người dùng)
    # =========================================================================
    ws_verif = wb.create_sheet(title="Verification")
    ws_verif.views.sheetView[0].showGridLines = True

    ws_verif["A1"] = "BẢNG ĐỐI CHIẾU TRUY XUẤT NGUỒN GỐC & TRẠNG THÁI XÁC MINH"
    ws_verif["A1"].font = title_font
    ws_verif["A2"] = "Lưu vết độc lập: Chuỗi OCR gốc, Gợi ý Vision AI, Giá trị đối soát và Bằng chứng nguồn"
    ws_verif["A2"].font = normal_font

    headers_verif = [
        "Mã chỉ tiêu",
        "Tên chỉ tiêu",
        "Kỳ báo cáo",
        "Phạm vi",
        "Chuỗi OCR thô (Raw OCR)",
        "Giá trị OCR trích xuất",
        "Gợi ý Vision AI",
        "Giá trị người dùng đối soát",
        "Trạng thái đối chiếu",
        "Trạng thái xác minh",
        "Lý do cần đối soát / Cảnh báo",
        "Trang nguồn PDF",
        "Dòng trích xuất gốc",
    ]
    ws_verif.append([])
    ws_verif.append(headers_verif)
    header_row_v = 4

    for col_idx in range(1, len(headers_verif) + 1):
        cell = ws_verif.cell(row=header_row_v, column=col_idx)
        cell.fill = header_fill_gray
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    current_row_v = header_row_v + 1
    for stmt in session.statements:
        for item in stmt.items:
            is_verified = item.status == "verified"

            # Xác định trạng thái so sánh giữa OCR và AI
            ocr_val = item.value
            ai_val = item.ai_suggestion
            if ai_val is None:
                compare_status = "Chưa gọi AI"
            elif ocr_val is not None and abs(ocr_val - ai_val) < 1.0:
                compare_status = "Trùng khớp"
            elif ocr_val is None:
                compare_status = "Chỉ có AI đọc được"
            else:
                compare_status = "Bất đồng"

            row_v = [
                item.code,
                item.name,
                item.period or stmt.period_label or "",
                item.scope,
                item.raw_ocr or "",
                ocr_val if ocr_val is not None else "",
                ai_val if ai_val is not None else "",
                item.user_value if item.user_value is not None else (item.value if is_verified else ""),
                compare_status,
                "Đã xác minh" if is_verified else "Cần đối soát",
                item.review_reason or "",
                item.page_number if item.page_number is not None else "",
                item.source_line or "",
            ]
            ws_verif.append(row_v)

            for c_idx in range(1, len(row_v) + 1):
                c = ws_verif.cell(row=current_row_v, column=c_idx)
                c.font = normal_font
                c.border = thin_border
                if c_idx in (6, 7, 8) and isinstance(row_v[c_idx - 1], (int, float)):
                    c.number_format = "#,##0"
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx == 10:
                    c.fill = fill_verified if is_verified else fill_needs_review
                    c.font = font_verified if is_verified else font_needs_review
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.alignment = Alignment(vertical="center")

            current_row_v += 1

    # =========================================================================
    # SHEET 3: Ratios (Tỷ số tài chính an toàn)
    # =========================================================================
    ws_ratios = wb.create_sheet(title="Ratios")
    ws_ratios.views.sheetView[0].showGridLines = True

    ws_ratios["A1"] = "BẢNG KẾT QUẢ TỶ SỐ TÀI CHÍNH"
    ws_ratios["A1"].font = title_font

    headers_ratios = [
        "Mã tỷ số",
        "Tên chỉ số tài chính",
        "Công thức chuẩn hóa",
        "Giá trị tính toán",
        "Đơn vị",
        "Trạng thái tính",
        "Căn cứ đầu vào & Ghi chú kiểm soát",
    ]
    ws_ratios.append([])
    ws_ratios.append(headers_ratios)
    header_row_r = 3

    for col_idx in range(1, len(headers_ratios) + 1):
        cell = ws_ratios.cell(row=header_row_r, column=col_idx)
        cell.fill = header_fill_blue
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ratio_definitions = [
        ("CURRENT_RATIO", "Current Ratio", "Tài sản ngắn hạn / Nợ ngắn hạn", "Lần"),
        ("DEBT_EQUITY", "Debt/Equity", "Tổng nợ phải trả / Vốn chủ sở hữu", "Lần"),
        ("ROA", "ROA (%)", "(Lợi nhuận sau thuế / Tổng tài sản bình quân) * 100", "%"),
        ("ROE", "ROE (%)", "(Lợi nhuận sau thuế / Vốn CSH bình quân) * 100", "%"),
    ]

    is_calculated = session.ratio_status == "calculated" and bool(session.ratios)
    current_row_r = header_row_r + 1

    for r_code, r_name, r_formula, r_unit in ratio_definitions:
        val = session.ratios.get(r_code) if is_calculated else None
        if is_calculated and val is not None:
            r_status = "Đã tính hợp lệ"
            r_note = "Các chỉ tiêu bắt buộc đã được xác minh cùng kỳ và cùng phạm vi."
        elif is_calculated and val is None:
            r_status = "N/A (Mẫu số bằng 0 hoặc thiếu số dư đầu kỳ)"
            r_note = "Công thức chia cho 0 hoặc không có số liệu đầu kỳ để tính bình quân."
        else:
            r_status = "Chưa thể tính"
            r_note = session.ratio_block_reason or "Còn chỉ tiêu bắt buộc chưa được xác minh (status: needs_review)."

        row_r = [
            r_code,
            r_name,
            r_formula,
            val if val is not None else "",
            r_unit,
            r_status,
            r_note,
        ]
        ws_ratios.append(row_r)

        for c_idx in range(1, len(row_r) + 1):
            c = ws_ratios.cell(row=current_row_r, column=c_idx)
            c.font = normal_font
            c.border = thin_border
            if c_idx == 4 and val is not None:
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.font = bold_font
            elif c_idx == 6:
                c.fill = fill_verified if (is_calculated and val is not None) else fill_needs_review
                c.font = font_verified if (is_calculated and val is not None) else font_needs_review
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(vertical="center")

        current_row_r += 1

    # =========================================================================
    # SHEET 4: Metadata (Truy xuất nguồn gốc & Quản lý phiên)
    # =========================================================================
    ws_meta = wb.create_sheet(title="Metadata")
    ws_meta.views.sheetView[0].showGridLines = True

    ws_meta["A1"] = "THÔNG TIN TRUY XUẤT NGUỒN GỐC & QUẢN LÝ PHIÊN ĐỐI SOÁT"
    ws_meta["A1"].font = title_font

    total_items_count = sum(len(s.items) for s in session.statements)
    verified_items_count = sum(
        1 for s in session.statements for it in s.items if it.status == "verified"
    )
    needs_review_count = total_items_count - verified_items_count

    meta_rows = [
        ("Mã phiên đối soát (Session ID)", session.session_id),
        ("Tên tệp BCTC gốc (Source File)", session.source_filename),
        ("Mã băm SHA-256 tệp gốc (File Hash)", session.source_file_hash or "Chưa ghi nhận"),
        ("Doanh nghiệp", session.company_name),
        ("Thời điểm khởi tạo phiên", session.created_at),
        ("Thời điểm cập nhật phiên", session.updated_at),
        ("Thời điểm xuất tệp Excel", datetime.now(timezone.utc).isoformat()),
        ("Phiên bản Schema Session", session.version),
        ("Hệ thống phát triển", "Financial_AI v0.7.0 (Giai đoạn 3)"),
        ("Tổng số chỉ tiêu trích xuất", total_items_count),
        ("Số chỉ tiêu ĐÃ XÁC MINH (Verified)", verified_items_count),
        ("Số chỉ tiêu CẦN ĐỐI SOÁT (Needs Review)", needs_review_count),
        ("Trạng thái tính tỷ số tài chính", session.ratio_status),
        ("Lý do chặn tính tỷ số (nếu có)", session.ratio_block_reason or "Không có (Đã tính thành công)"),
    ]

    ws_meta.append([])
    ws_meta.append(["Thuộc tính", "Giá trị ghi nhận"])
    ws_meta.cell(row=3, column=1).fill = header_fill_blue
    ws_meta.cell(row=3, column=1).font = header_font
    ws_meta.cell(row=3, column=2).fill = header_fill_blue
    ws_meta.cell(row=3, column=2).font = header_font

    for idx, (prop, val) in enumerate(meta_rows, start=4):
        ws_meta.append([prop, str(val)])
        c1 = ws_meta.cell(row=idx, column=1)
        c2 = ws_meta.cell(row=idx, column=2)
        c1.font = bold_font
        c1.border = thin_border
        c2.font = normal_font
        c2.border = thin_border

    # Tự động căn chỉnh độ rộng cột cho cả 4 sheet
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                # Tránh các dòng tiêu đề làm phình to cột
                if cell.row <= 2:
                    continue
                max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_reconciliation_to_json(session: ReconciliationSession) -> str:
    """
    Xuất phiên đối soát ra chuỗi JSON có cấu trúc đầy đủ, bảo toàn 100% metadata.
    Tuyệt đối không lưu khóa bí mật hoặc token.
    """
    return session.to_json(indent=2)


def load_reconciliation_from_json(json_str: str) -> ReconciliationSession:
    """
    Nạp phiên đối soát từ chuỗi JSON và kiểm tra schema.
    Quy tắc an toàn cốt lõi:
    - Bắt lỗi JSON sai định dạng hoặc thiếu trường bắt buộc.
    - Bảo toàn nguyên trạng 'needs_review', TUYỆT ĐỐI KHÔNG tự nâng bất kỳ chỉ tiêu nào lên 'verified'.
    """
    try:
        session = ReconciliationSession.from_json(json_str)
        return session
    except Exception as exc:
        raise ValueError(f"Không thể nạp phiên đối soát: JSON không đúng định dạng schema ({exc})") from exc


def save_session_to_processed_dir(
    session: ReconciliationSession,
    target_dir: Path | str = "data/processed",
) -> Path:
    """
    Lưu an toàn phiên đối soát vào thư mục processed bằng cơ chế Atomic Write
    (ghi ra tệp tạm cùng thư mục rồi đổi tên để tránh làm hỏng tệp khi lỗi giữa chừng).
    """
    dir_path = Path(target_dir)
    dir_path.mkdir(parents=True, exist_ok=True)

    # Khử ký tự đặc biệt trong tên tệp
    clean_filename = re.sub(r"[^\w\-.]", "_", session.source_filename)
    target_filename = f"session_{session.session_id}_{clean_filename}.json"
    target_filepath = dir_path / target_filename

    json_content = export_reconciliation_to_json(session)

    # Atomic write
    temp_file = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=dir_path,
            delete=False,
            suffix=".tmp",
        ) as tf:
            temp_file = Path(tf.name)
            tf.write(json_content)

        # Đổi tên tệp tạm thành tệp đích
        temp_file.replace(target_filepath)
        return target_filepath
    except Exception as exc:
        if temp_file and temp_file.exists():
            temp_file.unlink(missing_ok=True)
        raise IOError(f"Lỗi khi lưu phiên đối soát vào {target_filepath}: {exc}") from exc


def load_session_from_file(file_path: Path | str) -> ReconciliationSession:
    """Đọc tệp JSON từ đường dẫn và nạp thực thể ReconciliationSession."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Tệp phiên đối soát không tồn tại: {path}")

    with open(path, "r", encoding="utf-8") as f:
        content = f.read()

    return load_reconciliation_from_json(content)
